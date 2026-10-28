"""Read-only audit of missing/empty/error predictions and generation log failures."""
from collections import Counter
import hashlib
import json
from pathlib import Path
import re

ROOT=Path(__file__).resolve().parents[3]
BASE=ROOT/'Qwen_vl/outputs/audit_followup_20261008'
OUT=BASE/'next_runtime_repair_20261009/timeout_audit.json'
ERROR=re.compile(r'TimeoutError|timed out|timeout exceeded|CUDA out of memory|CUDA error:|Traceback \(most recent call last\)',re.I)
FAIL=re.compile(r'^\s*(?:FAILED\b|\[FAILED\]|Failed to obtain answer|CUDA out of memory|TimeoutError)',re.I)


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def main():
    question_file=ROOT/'LLaVA/playground/data/eval/textvqa/llava_textvqa_val_v051_ocr.jsonl'
    questions=[json.loads(l) for l in question_file.open()]
    keys=[(str(r['question_id']),r['text']) for r in questions]
    jobs=[]
    # Original completed v1.5 and NeXT TextVQA arms, with their registered logs.
    for name in ['lane1_v15','lane2_next']:
        plan=json.loads((BASE/f'rerun_batch/{name}_plan.json').read_text())
        for j in plan['jobs']:
            if j['task']=='textvqa':
                jobs.append(dict(label='original/'+j['key'],prediction=j['output'],log=j['generation_log']))
    # Current fully scored wait-only NeXT arms and the fixed beta=1 control.
    controls=list((BASE/'streamwait_continuation_20261008').glob('next_text_*.control.protocol.json'))
    controls+=list((BASE/'next_gap_diagnosis_20261008/stability').glob('next_text_K32_*_streamwait.control.protocol.json'))
    for p in controls:
        protocol=json.loads(p.read_text());paths=protocol['output_paths']
        jobs.append(dict(label='wait/'+p.name,prediction=paths['prediction'],log=paths['log'],
                         runtime=paths['runtime'],state=paths['state'],finished=paths['finished'],score=paths['score']))
    reports=[]
    for job in jobs:
        p=Path(job['prediction'])
        rows=[json.loads(l) for l in p.open()]
        actual=[(str(r['question_id']),r['prompt']) for r in rows]
        counts=Counter(actual)
        blanks=[i for i,r in enumerate(rows) if not isinstance(r.get('text'),str) or not r['text'].strip()]
        failures=[i for i,r in enumerate(rows) if isinstance(r.get('text'),str) and FAIL.search(r['text'])]
        report=dict(job,n=len(rows),expected_n=len(keys),ordered_keys_match=actual==keys,
                    missing=sum((Counter(keys)-counts).values()),unexpected=sum((counts-Counter(keys)).values()),
                    duplicate_composite_keys=sum(v-1 for v in counts.values() if v>1),
                    empty_or_nonstring=len(blanks),failure_marker_predictions=len(failures),
                    affected_row_positions=sorted(set(blanks+failures)),prediction_sha256=sha(p))
        log=Path(job['log'])
        if log.exists():
            text=log.read_text(errors='replace')
            report['generation_log_sha256']=sha(log)
            report['generation_log_error_matches']=dict(Counter(m.group(0) for m in ERROR.finditer(text)))
        else:report['generation_log_missing']=True
        if 'runtime' in job:
            report['runtime_rows']=sum(1 for _ in Path(job['runtime']).open())
            state=json.loads(Path(job['state']).read_text());marker=json.loads(Path(job['finished']).read_text())
            score=json.loads(Path(job['score']).read_text())
            report.update(generation_exit_code=state.get('generation_exit_code'),state_complete=state.get('complete'),
                          finished_success=marker.get('success'),score_prediction_sha_matches=score['prediction_sha256']==sha(p))
        reports.append(report)

    from openpyxl import load_workbook
    qwen=[]
    archived=Path('/media/disk2/YZX/research/audit_base_gap_20261003/archive_recovery/20261003_r2/old_predictions')
    qwen_sources=[('EADP','TextVQA',5000,archived/'old_TextVQA_VAL.xlsx'),
                  ('EADP','DocVQA',5349,archived/'old_DocVQA_VAL.xlsx')]
    current=ROOT/'Qwen_vl/outputs/stage1_roundtrip_pilot/legacy_full/acc'
    for arm in ['L_E_GATHER','L_R_MAIN025']:
        for task,n,file in [('AI2D',3088,'AI2D_TEST_pred_exact_matching_result.xlsx'),
                            ('HallusionBench',951,'HallusionBench_pred_auxmatch.xlsx')]:
            qwen_sources.append((arm,task,n,current/arm/file))
    for task,n in [('TextVQA_VAL',5000),('DocVQA_VAL',5349)]:
        qwen_sources.append(('L_R_MAIN025',task,n,current/'L_R_MAIN025'/f'{task}_pred_results.xlsx'))
    for arm,task,n,p in qwen_sources:
        book=load_workbook(p,read_only=True,data_only=True)
        records=book.active.iter_rows(values_only=True);header=next(records)
        pred_idx=header.index('prediction');id_idx=header.index('index')
        values=[];indices=[]
        for row in records:
            indices.append(str(row[id_idx]));values.append(row[pred_idx])
        blank=[i for i,v in enumerate(values) if v is None or isinstance(v,str) and not v.strip()]
        errors=[i for i,v in enumerate(values) if isinstance(v,str) and FAIL.search(v)]
        qwen.append(dict(arm=arm,task=task,path=str(p),sha256=sha(p),n=len(values),expected_n=n,
                         duplicate_indices=len(indices)-len(set(indices)),empty=len(blank),failure_markers=len(errors),
                         affected_indices=[indices[i] for i in sorted(set(blank+errors))],
                         scope='Prediction-cell completeness only; original runtime timeout settings/logs not independently revalidated here.'))
        book.close()
    report=dict(llava=reports,qwen=qwen,question_file_sha256=sha(question_file),
                interpretation='Missing/empty/error answers are distinct from nonempty wrong answers and from stream races. No GPU job launched or altered.')
    OUT.write_text(json.dumps(report,indent=2,ensure_ascii=False))
    for r in reports:
        print(json.dumps({k:r[k] for k in ['label','n','missing','empty_or_nonstring','failure_marker_predictions','generation_log_error_matches'] if k in r}))
    for r in qwen:
        print(json.dumps({k:r[k] for k in ['arm','task','n','empty','failure_markers','duplicate_indices']}))
    print('Saved '+str(OUT))


if __name__=='__main__':main()
